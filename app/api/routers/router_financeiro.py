# -*- coding: utf-8 -*-
"""
Router Financeiro - PMR + Pipeline de dados
Destino: app/api/routers/router_financeiro.py

CONTRATO DE RESPOSTA (rotas -filtrado / filtros / evolucao / notas / clientes-busca):
  retornam o objeto de dados DIRETO (sem envelope {"status":"success","data":...}),
  pois o FinanceiroArena.tsx consome r.data.<campo> diretamente. Em erro, devolvem
  estrutura vazia valida (nunca undefined) para o .map() do frontend nao quebrar.
"""
import traceback
import asyncio
import logging
from datetime import date
import pandas as pd
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query
from pydantic import BaseModel
from typing import Optional
from app.api.routers.router_auth import get_current_user

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1/financeiro", tags=["Financeiro - PMR"])


class PerguntaFinanceira(BaseModel):
    pergunta: str
    data_ini: date
    data_fim: date
    segmentos: Optional[str] = None
    regionais: Optional[str] = None
    status: Optional[str] = None
    motivos: Optional[str] = None
    cgc: Optional[str] = None
    historico: Optional[list[dict]] = None


@router.post("/agente/chat")
async def agente_chat(
    payload: PerguntaFinanceira,
    _: dict = Depends(get_current_user),
):
    pergunta = payload.pergunta.strip()
    if not pergunta:
        raise HTTPException(422, "A pergunta financeira nao pode ser vazia.")
    _validar_datas(payload.data_ini, payload.data_fim)
    try:
        from app.financeiro.agente_financeiro import construir_contexto, responder_pergunta
        segmento = _split(payload.segmentos)
        regional = _split(payload.regionais)
        contexto = await asyncio.to_thread(
            construir_contexto, payload.data_ini, payload.data_fim, segmento, regional, payload.cgc,
            _split(payload.status), _split(payload.motivos)
        )
        return {"resposta": await asyncio.to_thread(
            responder_pergunta, pergunta, contexto, payload.historico
        )}
    except RuntimeError as e:
        raise HTTPException(503, str(e))
    except Exception as e:
        logger.exception("agente financeiro: %s", e)
        raise HTTPException(500, "Erro ao consultar o agente financeiro.")


# ------------------------------------------------------------
# Estado do pipeline financeiro
# ------------------------------------------------------------

class FinanceiroState:
    pipeline_rodando:  bool = False
    is_recarga_total:  bool = False
    logs:              list = []


# ------------------------------------------------------------
# Lazy imports
# ------------------------------------------------------------

def _engine():
    from app.financeiro import pmr_engine
    return pmr_engine

def _store():
    from app.financeiro import s3_store
    return s3_store

def _pmp():
    from app.financeiro import pmp_engine
    return pmp_engine


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------

def _validar_datas(data_ini: date, data_fim: date) -> None:
    if data_ini > data_fim:
        raise HTTPException(422, "data_ini deve ser anterior a data_fim.")
    if (data_fim - data_ini).days > 365 * 5:
        raise HTTPException(422, "Janela maxima: 5 anos.")

def _so_admin(usuario: dict) -> None:
    if usuario.get("funcao") != "Administrador":
        raise HTTPException(403, "Apenas Administradores podem executar o pipeline financeiro.")

def _split(v):
    """Converte 'A,B,C' em ['A','B','C']; None/'' vira None."""
    if not v:
        return None
    itens = [x.strip() for x in v.split(",") if x.strip()]
    return itens or None


# ------------------------------------------------------------
# Status
# ------------------------------------------------------------

@router.get("/status")
async def status(_: dict = Depends(get_current_user)):
    def _check():
        store = _store()
        fontes = ["notas_saida", "contas_receber", "movimentacao_bancaria",
                  "notas_entrada", "contas_pagar"]
        return {
            "is_running":       FinanceiroState.pipeline_rodando,
            "is_recarga_total": FinanceiroState.is_recarga_total,
            "logs":             list(FinanceiroState.logs),
            "parquets": {
                fonte: [f"{a}/{m:02d}" for a, m in store.meses_disponiveis(fonte)]
                for fonte in fontes
            },
            "latest": {
                "fornecedores": "fornecedores/latest.parquet",
                "produtos": "sb1/latest.parquet",
                "lista_tecnica": "sg1/latest.parquet",
            },
            # Cadastros sem coluna de data não aparecem em ``meses_disponiveis``:
            # o indicador explícito evita que consumidores os tratem como
            # partições mensais.
            "snapshots": {
                "fornecedores": "latest",
                "produtos": "latest (SB1)",
                "lista_tecnica": "latest (SG1)",
            },
            "latest_only": ["fornecedores", "produtos", "lista_tecnica"],
        }
    return await asyncio.to_thread(_check)


@router.get("/pmp/global")
async def pmp_global(data_ini: date = Query(...), data_fim: date = Query(...),
                     e5_motbx: str = Query(None), d1_tp: str = Query(None),
                     fornecedor: str = Query(None), _: dict = Depends(get_current_user)):
    _validar_datas(data_ini, data_fim)
    return await asyncio.to_thread(_pmp().calcular_pmp_global, data_ini, data_fim,
                                   _split(e5_motbx), _split(d1_tp), _split(fornecedor))


@router.get("/pmp/fornecedores")
async def pmp_fornecedores(data_ini: date = Query(...), data_fim: date = Query(...),
                           limit: int = Query(100, ge=1, le=500),
                           offset: int = Query(0, ge=0),
                           e5_motbx: str = Query(None), d1_tp: str = Query(None),
                           fornecedor: str = Query(None),
                           _: dict = Depends(get_current_user)):
    _validar_datas(data_ini, data_fim)
    return await asyncio.to_thread(
        _pmp().calcular_pmp_fornecedores, data_ini, data_fim, limit, offset,
        _split(e5_motbx), _split(d1_tp), _split(fornecedor)
    )


def _pmp_filters(motivos, tipos, fornecedores, clifor):
    return (_split(motivos), _split(tipos), _split(fornecedores), _split(clifor))


@router.get("/pmp/filtros")
async def pmp_filtros(data_ini: date = Query(...), data_fim: date = Query(...),
                      _: dict = Depends(get_current_user)):
    def listar():
        base = _pmp()._base(data_ini, data_fim)
        if base.empty:
            return {"e5_motbx": [], "d1_tp": [], "fornecedor": [], "motivos": [], "tipos": [], "fornecedores": [], "clifor": []}
        valores = lambda col: sorted(_pmp()._norm_key(_pmp()._col(base, col)).replace("", "SEM VALOR").unique().tolist())
        fornecedores = valores("clifor")
        motivos = valores("e5_motbx")
        tipos = valores("d1_tp")
        return {"e5_motbx": motivos, "d1_tp": tipos, "fornecedor": fornecedores,
                "motivos": motivos, "tipos": tipos, "fornecedores": fornecedores, "clifor": fornecedores}
    return await asyncio.to_thread(listar)


@router.get("/pmp/global-filtrado")
async def pmp_global_filtrado(data_ini: date = Query(...), data_fim: date = Query(...),
                              motivos: str = Query(None), tipos: str = Query(None),
                              fornecedores: str = Query(None), clifor: str = Query(None),
                              e5_motbx: str = Query(None), d1_tp: str = Query(None),
                              fornecedor: str = Query(None),
                              _: dict = Depends(get_current_user)):
    _validar_datas(data_ini, data_fim)
    f = _pmp_filters(motivos or e5_motbx, tipos or d1_tp, fornecedores or fornecedor, clifor)
    return await asyncio.to_thread(_pmp().calcular_pmp_global, data_ini, data_fim, *f)


@router.get("/pmp/fornecedores-filtrado")
async def pmp_fornecedores_filtrado(data_ini: date = Query(...), data_fim: date = Query(...),
                                    motivos: str = Query(None), tipos: str = Query(None),
                                    fornecedores: str = Query(None), clifor: str = Query(None),
                                    e5_motbx: str = Query(None), d1_tp: str = Query(None),
                                    fornecedor: str = Query(None),
                                    limit: int = Query(100, ge=1, le=500),
                                    offset: int = Query(0, ge=0),
                                    _: dict = Depends(get_current_user)):
    _validar_datas(data_ini, data_fim)
    f = _pmp_filters(motivos or e5_motbx, tipos or d1_tp, fornecedores or fornecedor, clifor)
    return await asyncio.to_thread(_pmp().calcular_pmp_fornecedores, data_ini, data_fim, limit, offset, *f)


@router.get("/pmp/resumo")
async def pmp_resumo(data_ini: date = Query(...), data_fim: date = Query(...),
                     motivos: str = Query(None), tipos: str = Query(None),
                     fornecedores: str = Query(None), clifor: str = Query(None),
                     e5_motbx: str = Query(None), d1_tp: str = Query(None),
                     fornecedor: str = Query(None),
                     _: dict = Depends(get_current_user)):
    _validar_datas(data_ini, data_fim)
    f = _pmp_filters(motivos or e5_motbx, tipos or d1_tp, fornecedores or fornecedor, clifor)
    return await asyncio.to_thread(_pmp().calcular_pmp_resumo, data_ini, data_fim, *f)


@router.get("/pmp/exportar")
async def pmp_exportar(data_ini: date = Query(...), data_fim: date = Query(...),
                       motivos: str = Query(None), tipos: str = Query(None),
                       fornecedores: str = Query(None), clifor: str = Query(None),
                       e5_motbx: str = Query(None), d1_tp: str = Query(None),
                       fornecedor: str = Query(None),
                       _: dict = Depends(get_current_user)):
    import io
    from fastapi.responses import StreamingResponse
    _validar_datas(data_ini, data_fim)
    f = _pmp_filters(motivos or e5_motbx, tipos or d1_tp, fornecedores or fornecedor, clifor)

    def gerar():
        import pandas as pd
        base = _pmp()._base(data_ini, data_fim, *f)
        fornecedores_data = _pmp().calcular_pmp_fornecedores(data_ini, data_fim, 100000, 0, *f)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            # O bruto permanece no grão E5; abas separadas facilitam a
            # reconciliação com SE2, SF1 e SD1 sem alterar o PMP.
            base.to_excel(writer, index=False, sheet_name="Pagamentos E5")
            for prefix, sheet in (("e2_", "Titulos SE2"), ("f1_", "Notas SF1"),
                                  ("d1_", "Itens SD1")):
                cols = [c for c in base.columns if c.startswith(prefix)]
                if cols:
                    base[cols].to_excel(writer, index=False, sheet_name=sheet)
            pd.DataFrame(fornecedores_data["fornecedores"]).to_excel(
                writer, index=False, sheet_name="Fornecedores")
            narrativa = [
                ("METODOLOGIA DO PMP — PRAZO MÉDIO DE PAGAMENTO",),
                ("Imagine a história de uma compra: a nota fiscal nasce na SF1, seus itens são detalhados na SD1, o título é registrado na SE2 e cada baixa aparece na SE5. O PMP acompanha essa história até o dinheiro sair da empresa.",),
                ("A origem do cálculo é sempre a SE5: começamos em cada movimento bancário efetivamente baixado e, a partir dele, voltamos no tempo para localizar o título correspondente na SE2. Por isso, pagamentos parciais continuam separados e o valor de cada movimento E5 é o peso da média.",),
                ("O cruzamento principal é SE5 → SE2 pela chave do título e pelo fornecedor (CLIFOR). Só entram no PMP os movimentos SE5 que encontram um título de fornecedor na SE2. E5 é uma tabela mista: também possui recebimentos e outros movimentos, que não devem ser tratados como pagamento a fornecedor.",),
                ("A data de pagamento nasce na E5, em e5_data. Para medir o prazo, exigimos a data de emissão do título na SE2. A SF1 fornece a nota de entrada e seu valor bruto; a SD1 fornece o detalhe e o tipo de compra. Nenhuma delas substitui a emissão do título. Assim, o PMP Pagamento é a diferença entre a emissão da SE2 e a baixa na E5.",),
                ("O PMP Vencimento compara a emissão com e2_vencrea, que representa o vencimento real do título. O PMP Condição de Pagamento compara a emissão com e2_vencto, que representa o prazo contratual originalmente informado.",),
                ("A fórmula é: PMP = SOMA(dias de cada movimento × valor do movimento E5) ÷ SOMA(valor dos movimentos E5). O valor maior pesa mais, porque representa mais caixa efetivamente pago.",),
                ("O Delta de Atraso é: PMP Pagamento − PMP Condição de Pagamento. Resultado positivo indica pagamento depois do prazo contratado; resultado negativo indica pagamento antes do prazo.",),
                ("Depois de partir da E5, a SF1 é consultada no nível da nota fiscal para localizar f1_valbrut, o valor bruto da entrada. A SD1 está no nível do item e serve para explicar a composição do produto e o tipo de compra, não para multiplicar o pagamento.",),
                ("O valor bruto da SF1 pode aparecer repetido no detalhe quando uma nota possui vários movimentos bancários na SE5. Isso é esperado: não somamos f1_valbrut novamente para calcular o PMP. O peso do cálculo é o valor de cada movimento E5.",),
                ("Também não deduplicamos movimentos E5 por chave do título. Duas baixas para a mesma parcela podem representar pagamento parcial, complemento, juros ou outra movimentação legítima. Cada linha deve permanecer auditável.",),
                ("Exemplo real deste arquivo: " + (
                    "o fornecedor " + str(base.iloc[0].get("d1_fornecedor", base.iloc[0].get("clifor", "não identificado"))) +
                    " teve o título " + str(base.iloc[0].get("e2_num", "não identificado")) +
                    ", parcela " + str(base.iloc[0].get("e2_parcela", "única")) +
                    ", com movimento E5 de R$ " + f"{float(base.iloc[0].get('e5_valor', 0) or 0):,.2f}" +
                    ", emitido em " + str(pd.to_datetime(base.iloc[0].get("e2_emissao", "")).strftime("%d/%m/%Y")) +
                    " e pago em " + str(pd.to_datetime(base.iloc[0].get("e5_data", "")).strftime("%d/%m/%Y")) +
                    ". Portanto, este movimento contribui com o prazo entre essas duas datas."
                    if not base.empty else "não há movimentos no período para montar um exemplo."
                ),),
                ("Quando o Excel é filtrado por período, motivo, tipo ou fornecedor, o recálculo usa somente as linhas E5 que permanecem no filtro. Por isso, o resultado pode mudar sem que exista erro: mudou o conjunto de pagamentos e, consequentemente, os pesos da média.",),
                ("Fontes Protheus utilizadas: SE5 (movimentos e baixas bancárias), SE2 (títulos a pagar e vencimentos), SF1 (notas fiscais de entrada) e SD1 (itens das notas).",),
                ("A aba Pagamentos E5 é a base auditável do cálculo; Titulos SE2, Notas SF1 e Itens SD1 ajudam a reconstituir a história de cada pagamento.",),
            ]
            pd.DataFrame({"Metodologia": [linha[0] for linha in narrativa]}).to_excel(
                writer, index=False, sheet_name="Metodologia"
            )
        output.seek(0)
        return output
    output = await asyncio.to_thread(gerar)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="PMP_{data_ini}_{data_fim}.xlsx"'})


# ------------------------------------------------------------
# PMR Global / Regional / Clientes (rotas autenticadas — inalteradas)
# ------------------------------------------------------------

@router.get("/pmr/global")
async def pmr_global(
    data_ini: date = Query(..., description="Inicio da janela (f2_emissao)"),
    data_fim: date = Query(..., description="Fim da janela (f2_emissao)"),
    _: dict = Depends(get_current_user),
):
    _validar_datas(data_ini, data_fim)
    try:
        return await asyncio.to_thread(_engine().calcular_pmr_global, data_ini, data_fim)
    except Exception as e:
        logger.exception("pmr_global: %s", e)
        raise HTTPException(500, f"Erro ao calcular PMR global: {e}")


@router.get("/pmr/regional")
async def pmr_regional(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    _: dict = Depends(get_current_user),
):
    _validar_datas(data_ini, data_fim)
    try:
        return await asyncio.to_thread(_engine().calcular_pmr_regional, data_ini, data_fim)
    except Exception as e:
        logger.exception("pmr_regional: %s", e)
        raise HTTPException(500, f"Erro ao calcular PMR regional: {e}")


@router.get("/pmr/clientes")
async def pmr_clientes(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    limit:    int  = Query(50, ge=1, le=500),
    offset:   int  = Query(0,  ge=0),
    _: dict = Depends(get_current_user),
):
    _validar_datas(data_ini, data_fim)
    try:
        return await asyncio.to_thread(
            _engine().calcular_pmr_clientes, data_ini, data_fim, None, None, limit, offset
        )
    except Exception as e:
        logger.exception("pmr_clientes: %s", e)
        raise HTTPException(500, f"Erro ao calcular PMR por cliente: {e}")


# ------------------------------------------------------------
# Pipeline - background tasks
# ------------------------------------------------------------

def _run_pipeline(recarga_total: bool) -> None:
    from app.financeiro.pipeline_financeiro import executar_pipeline

    FinanceiroState.pipeline_rodando  = True
    FinanceiroState.is_recarga_total  = recarga_total
    FinanceiroState.logs              = []

    def _log(msg: str) -> None:
        FinanceiroState.logs.append(msg)
        logger.info("[financeiro] %s", msg)

    try:
        executar_pipeline(recarga_total=recarga_total, log_callback=_log)
    except Exception as e:
        _log(f"ERRO inesperado no pipeline: {e}")
        logger.exception("pipeline_financeiro falhou: %s", e)
    finally:
        _engine().limpar_cache()
        FinanceiroState.pipeline_rodando  = False
        FinanceiroState.is_recarga_total  = False


@router.post("/pipeline")
async def pipeline_atualizar(
    background_tasks: BackgroundTasks,
    usuario: dict = Depends(get_current_user),
):
    _so_admin(usuario)
    if FinanceiroState.pipeline_rodando:
        raise HTTPException(409, "Pipeline financeiro ja esta em execucao.")
    background_tasks.add_task(asyncio.to_thread, _run_pipeline, False)
    return {"status": "iniciado", "modo": "rolling_3_meses"}


@router.post("/pipeline/recarga")
async def pipeline_recarga(
    background_tasks: BackgroundTasks,
    usuario: dict = Depends(get_current_user),
):
    _so_admin(usuario)
    if FinanceiroState.pipeline_rodando:
        raise HTTPException(409, "Pipeline financeiro ja esta em execucao.")
    background_tasks.add_task(asyncio.to_thread, _run_pipeline, True)
    return {"status": "iniciado", "modo": "recarga_total_desde_2023"}


# ------------------------------------------------------------
# Filtros disponiveis (regionais e segmentos do periodo)
# ------------------------------------------------------------

@router.get("/pmr/filtros")
async def get_pmr_filtros(
    data_ini: date = Query(..., description="Data de início (YYYY-MM-DD)"),
    data_fim: date = Query(..., description="Data de fim (YYYY-MM-DD)")
):
    try:
        filtros = await asyncio.to_thread(_engine().listar_filtros, data_ini, data_fim)
        return filtros  # {"regionais": [...], "segmentos": [...]}
    except Exception as e:
        print("\n" + "="*50)
        print("ERRO DETECTADO NA ROTA DE FILTROS")
        traceback.print_exc()
        print("="*50 + "\n")
        return {"regionais": [], "segmentos": []}


@router.get("/pmr/clientes-busca")
async def get_clientes_busca(
    data_ini: date = Query(..., description="Data de início (YYYY-MM-DD)"),
    data_fim: date = Query(..., description="Data de fim (YYYY-MM-DD)"),
    q: str = Query("", description="Texto de busca (nome ou CNPJ)"),
    limit: int = Query(20, ge=1, le=100),
):
    """Autocomplete do filtro de cliente."""
    try:
        clientes = await asyncio.to_thread(
            _engine().listar_clientes_busca, data_ini, data_fim, q, limit
        )
        return {"clientes": clientes}
    except Exception as e:
        logger.exception("clientes-busca: %s", e)
        return {"clientes": []}


# ------------------------------------------------------------
# Rotas filtradas (multi-valor: regionais/segmentos comma-separated)
# ------------------------------------------------------------

@router.get("/pmr/global-filtrado")
async def get_pmr_global_filtrado(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    regional:  str = Query(None, description="Uma regional (compat)"),
    segmento:  str = Query(None, description="Um segmento (compat)"),
    regionais: str = Query(None, description="Multiplas regionais separadas por virgula"),
    segmentos: str = Query(None, description="Multiplos segmentos separados por virgula"),
    status: str = Query(None, description="Status do cliente"),
    motivos: str = Query(None, description="Multiplos motivos da E5"),
    cgc:       str = Query(None, description="Filtro opcional por CNPJ do cliente"),
):
    try:
        reg = _split(regionais) or regional
        seg = _split(segmentos) or segmento
        stat = _split(status) or status
        mot = _split(motivos) or motivos
        dados = await asyncio.to_thread(
            _engine().calcular_pmr_global, data_ini, data_fim, seg, reg, cgc, stat, mot
        )
        return dados
    except Exception as e:
        logger.exception("global-filtrado: %s", e)
        return {"periodo": {"data_ini": str(data_ini), "data_fim": str(data_fim),
                            "dias_periodo": max((data_fim - data_ini).days, 1)},
                "notas_base": 0, "notas_pagas": 0, "valor_total": 0.0, "valor_por_dia": 0.0,
                "pmr_pagamento": None, "pmr_vencimento": None, "pmr_cond_pag": None,
                "delta_atraso": None}


@router.get("/pmr/regional-filtrado")
async def get_pmr_regional_filtrado(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    segmento:  str = Query(None),
    segmentos: str = Query(None, description="Multiplos segmentos separados por virgula"),
    regional:   str = Query(None, description="Uma regional (compat)"),
    regionais:  str = Query(None, description="Multiplas regionais separadas por virgula"),
    status:     str = Query(None, description="Status do cliente"),
    motivos:    str = Query(None, description="Multiplos motivos da E5"),
):
    try:
        seg = _split(segmentos) or segmento
        reg = _split(regionais) or regional
        stat = _split(status) or status
        mot = _split(motivos) or motivos
        dados = await asyncio.to_thread(
            _engine().calcular_pmr_regional, data_ini, data_fim, seg, reg, stat, mot
        )
        return dados  # {"periodo": {...}, "regionais": [...]}
    except Exception as e:
        logger.exception("regional-filtrado: %s", e)
        return {"periodo": {"data_ini": str(data_ini), "data_fim": str(data_fim)}, "regionais": []}


@router.get("/pmr/clientes-filtrado")
async def get_pmr_clientes_filtrado(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    regional:  str = Query(None),
    segmento:  str = Query(None),
    regionais: str = Query(None, description="Multiplas regionais separadas por virgula"),
    segmentos: str = Query(None, description="Multiplos segmentos separados por virgula"),
    status: str = Query(None, description="Status do cliente"),
    motivos: str = Query(None, description="Multiplos motivos da E5"),
    limit: int = Query(50),
    offset: int = Query(0),
):
    try:
        reg = _split(regionais) or regional
        seg = _split(segmentos) or segmento
        stat = _split(status) or status
        mot = _split(motivos) or motivos
        dados = await asyncio.to_thread(
            _engine().calcular_pmr_clientes, data_ini, data_fim, seg, reg, limit, offset, stat, mot
        )
        return dados  # {"periodo": {...}, "total": N, "clientes": [...]}
    except Exception as e:
        logger.exception("clientes-filtrado: %s", e)
        return {"periodo": {"data_ini": str(data_ini), "data_fim": str(data_fim)},
                "total": 0, "clientes": []}


@router.get("/pmr/evolucao")
async def get_pmr_evolucao(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    regional:  str = Query(None),
    segmento:  str = Query(None),
    regionais: str = Query(None, description="Multiplas regionais separadas por virgula"),
    segmentos: str = Query(None, description="Multiplos segmentos separados por virgula"),
    cgc:       str = Query(None, description="Filtro opcional por CNPJ do cliente"),
    status:    str = Query(None, description="Status do cliente"),
    motivos:   str = Query(None, description="Multiplos motivos da E5"),
):
    """Evolucao mes a mes, agrupada pelo mes de recebimento (e5_data)."""
    try:
        reg = _split(regionais) or regional
        seg = _split(segmentos) or segmento
        stat = _split(status) or status
        mot = _split(motivos) or motivos
        dados = await asyncio.to_thread(
            _engine().calcular_pmr_mensal, data_ini, data_fim, seg, reg, cgc, stat, mot
        )
        return dados  # {"periodo": {...}, "meses": [...]}
    except Exception as e:
        logger.exception("evolucao: %s", e)
        return {"periodo": {"data_ini": str(data_ini), "data_fim": str(data_fim)}, "meses": []}


@router.get("/pmr/notas")
async def get_pmr_notas(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    cgc:      str  = Query(None, description="CNPJ do cliente (compatibilidade)"),
    razao_social: str = Query(None, description="Razao social consolidada"),
    segmentos: str = Query(None, description="Segmentos aplicados ao detalhe"),
    regionais: str = Query(None, description="Regionais aplicadas ao detalhe"),
    status: str = Query(None, description="Status aplicado ao detalhe"),
    motivos: str = Query(None, description="Motivos E5 aplicados ao detalhe"),
):
    """Drill-down: parcelas de todos os CNPJs de uma razao social."""
    try:
        if not razao_social and not cgc:
            raise HTTPException(422, "Informe razao_social ou cgc.")
        if razao_social:
            dados = await asyncio.to_thread(
                _engine().obter_notas_razao_social, data_ini, data_fim,
                razao_social, _split(segmentos), _split(regionais), _split(status), _split(motivos)
            )
        else:
            dados = await asyncio.to_thread(_engine().obter_notas_cliente, data_ini, data_fim, cgc)
        return dados  # {"cgc", "nome", "total", "resumo", "notas": [...]}
    except Exception as e:
        logger.exception("notas: %s", e)
        return {"cgc": cgc, "nome": None, "total": 0, "resumo": {}, "notas": []}


# ------------------------------------------------------------
# Exportacao Excel — 5 abas (inclui Evolucao Mensal)
# ------------------------------------------------------------

@router.get("/pmr/exportar")
async def pmr_exportar(
    data_ini: date = Query(...),
    data_fim: date = Query(...),
    segmento: str = Query(None),
    regional: str = Query(None),
    segmentos: str = Query(None),
    regionais: str = Query(None),
    status: str = Query(None, description="Status do cliente"),
    motivos: str = Query(None, description="Multiplos motivos da E5"),
    _: dict = Depends(get_current_user),
):
    """
    Gera um Excel com 5 abas para validacao e para servir de base ao agente de IA:
      1. Dados Brutos     - linha a linha (nivel parcela) com todas as variaveis
      2. Por Regional     - PMR agregado por regional
      3. Resumo por Cliente - PMR agregado por razao social
      4. Evolucao Mensal  - PMR mes a mes por mes de pagamento (+ flag maturacao)
      5. Metodologia      - explicacao + exemplos numericos provados
    """
    import io
    from fastapi.responses import StreamingResponse

    _validar_datas(data_ini, data_fim)

    def _gerar():
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        engine = _engine()
        reg = _split(regionais) or regional
        seg = _split(segmentos) or segmento
        stat = _split(status) or status
        mot = _split(motivos) or motivos

        df_bruto    = engine.obter_dados_brutos(data_ini, data_fim, seg, reg, status=stat, motivo=mot)
        d_regional  = engine.calcular_pmr_regional(data_ini, data_fim, seg, reg, stat, mot)
        d_clientes  = engine.calcular_pmr_clientes(data_ini, data_fim, seg, reg, 100000, 0, stat, mot)
        d_global    = engine.calcular_pmr_global(data_ini, data_fim, seg, reg, status=stat, motivo=mot)
        d_mensal    = engine.calcular_pmr_mensal(data_ini, data_fim, seg, reg, status=stat, motivo=mot)

        wb = openpyxl.Workbook()
        hdr_fill = PatternFill("solid", fgColor="1E3A5F")
        hdr_font = Font(bold=True, color="FFFFFF", size=10)

        # ABA 1: Dados Brutos
        ws1 = wb.active
        ws1.title = "Dados Brutos"
        if not df_bruto.empty:
            for j, col in enumerate(df_bruto.columns, 1):
                c = ws1.cell(row=1, column=j, value=col)
                c.fill, c.font = hdr_fill, hdr_font
                c.alignment = Alignment(horizontal="center")
            for i, row in enumerate(df_bruto.itertuples(index=False), 2):
                for j, v in enumerate(row, 1):
                    ws1.cell(row=i, column=j, value=v)

        # ABA 2: Por Regional
        ws2 = wb.create_sheet("Por Regional")
        hdrs2 = ["Regional", "Movimentos E5", "Valor Recebido E5 (R$)",
                 "PMR Pagamento (dias)", "PMR Vencimento (dias)", "PMR Cond.Pag. (dias)", "Delta Atraso (dias)"]
        for j, h in enumerate(hdrs2, 1):
            c = ws2.cell(row=1, column=j, value=h); c.fill, c.font = hdr_fill, hdr_font
        for i, r in enumerate(d_regional.get("regionais", []), 2):
            ws2.cell(row=i, column=1, value=r["regional"])
            ws2.cell(row=i, column=2, value=r["notas_pagas"])
            ws2.cell(row=i, column=3, value=r["valor_total"])
            ws2.cell(row=i, column=4, value=r["pmr_pagamento"])
            ws2.cell(row=i, column=5, value=r["pmr_vencimento"])
            ws2.cell(row=i, column=6, value=r["pmr_cond_pag"])
            ws2.cell(row=i, column=7, value=r["delta_atraso"])

        # ABA 3: Resumo por Cliente
        ws3 = wb.create_sheet("Resumo por Cliente")
        hdrs3 = ["Razão Social", "CNPJs", "Códigos Cliente", "Regional", "Segmento", "Movimentos E5", "Valor Recebido E5 (R$)",
                 "PMR Pagamento (dias)", "PMR Vencimento (dias)", "PMR Cond. Pagamento (dias)", "Delta Atraso (dias)"]
        for j, h in enumerate(hdrs3, 1):
            c = ws3.cell(row=1, column=j, value=h); c.fill, c.font = hdr_fill, hdr_font
        for i, r in enumerate(d_clientes.get("clientes", []), 2):
            ws3.cell(row=i, column=1, value=r["nome"])
            ws3.cell(row=i, column=2, value=", ".join(r.get("cnpjs", [])))
            ws3.cell(row=i, column=3, value=", ".join(r.get("codigos_cliente", [])))
            ws3.cell(row=i, column=4, value=r["regional"])
            ws3.cell(row=i, column=5, value=r["segmento"])
            ws3.cell(row=i, column=6, value=r["notas_pagas"])
            ws3.cell(row=i, column=7, value=r["valor_total"])
            ws3.cell(row=i, column=8, value=r["pmr_pagamento"])
            ws3.cell(row=i, column=9, value=r["pmr_vencimento"])
            ws3.cell(row=i, column=10, value=r["pmr_cond_pag"])
            ws3.cell(row=i, column=11, value=r["delta_atraso"])

        # ABA 4: Evolucao Mensal
        ws5 = wb.create_sheet("Evolução Mensal")
        hdrs5 = ["Mês de Recebimento E5", "Movimentos E5", "Valor Recebido E5 (R$)", "Valor Recebido (R$)",
                 "PMR Pagamento (dias)", "PMR Vencimento (dias)", "PMR Cond. Pagamento (dias)",
                 "Delta Atraso (dias)", "Em Maturação?"]
        for j, h in enumerate(hdrs5, 1):
            c = ws5.cell(row=1, column=j, value=h); c.fill, c.font = hdr_fill, hdr_font
        for i, m in enumerate(d_mensal.get("meses", []), 2):
            ws5.cell(row=i, column=1, value=m["mes"])
            ws5.cell(row=i, column=2, value=m["notas_pagas"])
            ws5.cell(row=i, column=3, value=m["valor_total"])
            ws5.cell(row=i, column=4, value=m["valor_recebido"])
            ws5.cell(row=i, column=5, value=m["pmr_pagamento"])
            ws5.cell(row=i, column=6, value=m["pmr_vencimento"])
            ws5.cell(row=i, column=7, value=m["pmr_cond_pag"])
            ws5.cell(row=i, column=8, value=m["delta_atraso"])
            ws5.cell(row=i, column=9, value="SIM" if m.get("em_maturacao") else "nao")

        # ABA 5: Metodologia (com exemplos numericos provados)
        ws4 = wb.create_sheet("Metodologia")
        titulo_font = Font(bold=True, size=12, color="1E3A5F")

        vt   = d_global.get("valor_total", 0) or 0
        dias = d_global.get("periodo", {}).get("dias_periodo", 1) or 1
        vpd  = d_global.get("valor_por_dia", 0) or 0
        pmr_pag = d_global.get("pmr_pagamento") or 0

        exemplo = ""
        if not df_bruto.empty:
            r = df_bruto.iloc[0]
            emissao = pd.to_datetime(r.get("Data Emissao"), errors="coerce")
            pagamento = pd.to_datetime(r.get("Data Pagamento (E5)"), errors="coerce")
            dias_exemplo = (pagamento - emissao).days if pd.notna(emissao) and pd.notna(pagamento) else None
            exemplo = (
                f"O cliente {r.get('Razao Social', 'não identificado')} tinha o título "
                f"{r.get('Num. Titulo', 'não identificado')}, no valor de R$ "
                f"{float(r.get('Valor Titulo E1 (R$)', 0) or 0):,.2f}, parcela "
                f"{r.get('Parcela', 'não identificada')}, emitido em "
                f"{emissao.strftime('%d/%m/%Y') if pd.notna(emissao) else 'data não disponível'} "
                f"e pagou em {pagamento.strftime('%d/%m/%Y') if pd.notna(pagamento) else 'data não disponível'}; "
                f"portanto, o PMR deste movimento é de {dias_exemplo if dias_exemplo is not None else 'não calculável'} dias."
            )

        texto = [
            ("CÁLCULO DO PMR — PRAZO MÉDIO DE RECEBIMENTO", True),
            ("", False),
            ("A história começa sempre na E5: cada movimento bancário informa que um recebimento aconteceu. A partir dessa baixa, voltamos no tempo para procurar o título correspondente na SE1 e, depois, a nota SF2/F2 que originou aquele título.", False),
            ("Somente depois de encontrar o título na SE1 e a nota correspondente na SF2/F2, com sua data de emissão, calculamos o prazo. A média é ponderada pelo valor efetivamente recebido em cada movimento E5.", False),
            ("", False),
            ("EXEMPLO COM DADOS DESTE ARQUIVO:", True),
            (exemplo or "Não há pagamentos no período para montar um exemplo individual.", False),
            ("", False),
            ("FÓRMULA GERAL:", True),
            ("PMR = SOMA(dias_i × e5_valor_i) ÷ SOMA(e5_valor_i).", False),
            ("A origem continua sendo a E5. Para cada e5_data, buscamos primeiro o título e os vencimentos na SE1 e depois a emissão e o valor bruto da SF2/F2. Então calculamos: PMR Pagamento = e5_data − f2_emissao; PMR Vencimento = e1_vencrea − f2_emissao; PMR Condição de Pagamento = e1_vencto − f2_emissao.", False),
            ("", False),
            ("COMO AS TABELAS SE ENCAIXAM:", True),
            ("A E5 informa cada entrada de dinheiro e a data em que ela ocorreu; ela é a tabela de origem do cálculo. A partir de cada linha E5, buscamos a parcela e os vencimentos na SE1 e, em seguida, a nota SF2/F2 e sua emissão e valor bruto.", False),
            ("O caminho de reconciliação é E5 → SE1 → SF2. Somente movimentos E5 que encontram uma parcela SE1 e uma nota SF2/F2 com data de emissão conseguem gerar PMR. Sem a emissão da SF2, não existe uma origem temporal confiável para o cálculo.", False),
            ("A aba Dados Brutos fica no nível da parcela/movimento E5. Uma NF pode ter várias parcelas e uma parcela pode ter vários movimentos bancários; cada movimento permanece visível para auditoria.", False),
            ("O valor bruto da F2/SF2 pode aparecer repetido no detalhe porque uma mesma nota pode possuir vários movimentos E5. Isso não significa que o faturamento foi multiplicado: para faturamento, a NF é contada uma vez; para PMR, o peso é o valor de cada pagamento E5.", False),
            ("", False),
            ("REGRAS DE INCLUSÃO:", True),
            ("Somente movimentos E5 ligados a títulos SE1 e notas SF2/F2 com data de emissão entram no cálculo. Notas ainda não recebidas não entram no PMR, pois ainda não existe uma linha de origem na E5 com data de pagamento.", False),
            ("A consolidação do resumo é feita pela Razão Social, mas o detalhe preserva CNPJ, loja, parcela e movimento bancário. Assim, uma empresa com vários CNPJs aparece consolidada no resumo sem perder a rastreabilidade.", False),
            ("Filtros aplicados ao Excel alteram o conjunto de movimentos E5 e todos os PMRs são recalculados sobre o novo conjunto.", False),
            ("", False),
            ("EXEMPLO DE RECÁLCULO POR FILTRO:", True),
            ("Imagine três recebimentos: R$ 100.000 pagos em 40 dias, R$ 50.000 em 50 dias e R$ 10.000 em 65 dias.", False),
            ("Sem filtro: (40 × 100.000 + 50 × 50.000 + 65 × 10.000) ÷ 160.000 = 44,7 dias.", False),
            ("Se o filtro retirar o terceiro movimento: (40 × 100.000 + 50 × 50.000) ÷ 150.000 = 43,3 dias.", False),
            ("O mesmo raciocínio vale para PMR Pagamento, PMR Vencimento, PMR Condição de Pagamento e Delta.", False),
            ("", False),
            ("EXEMPLO NUMÉRICO:", True),
            ("Uma parcela emitida em 01/03 e paga em 10/04 gera 40 dias. Se o recebimento foi de R$ 100.000, sua contribuição é 40 × 100.000 = 4.000.000.", False),
            ("Outra parcela paga em 50 dias e no valor de R$ 50.000 contribui com 2.500.000; uma terceira, paga em 65 dias e no valor de R$ 10.000, contribui com 650.000.", False),
            ("A soma dos produtos dias × valor é 7.150.000 e a soma dos valores é 160.000; logo, o PMR é 44,7 dias, e não a média simples de 51,7 dias.", False),
            ("A parcela de maior valor exerce maior influência no resultado. Essa é a razão da ponderação por valor.", False),
            ("", False),
            ("IMPACTO DE UM DIA DE PMR NO CAPITAL DE GIRO:", True),
            ("Fórmula: valor por dia = valor recebido no período ÷ quantidade de dias do período.", False),
            (f"Neste recorte, o valor recebido é R$ {vt:,.2f}, em {dias} dias; portanto, cada dia representa aproximadamente R$ {vpd:,.2f}.", False),
            ("Reduzir o PMR em um dia libera, em uma estimativa de ordem de grandeza, esse valor de caixa preso no contas a receber.", False),
            (f"Exemplo: reduzir o PMR de {pmr_pag:.0f} para {max(pmr_pag-5,0):.0f} dias representaria aproximadamente R$ {vpd*5:,.2f}.", False),
            ("Essa estimativa pressupõe recebimentos aproximadamente uniformes e não substitui uma apuração contábil.", False),
            ("", False),
            ("EVOLUÇÃO MENSAL:", True),
            ("A aba Evolução Mensal agrupa os movimentos pelo mês de recebimento, usando e5_data, mas calcula cada prazo sempre a partir da emissão da SF2.", False),
            ("Os dois últimos meses podem estar em maturação: pagamentos mais lentos ainda não aconteceram e, por isso, o PMR pode parecer artificialmente menor.", False),
            ("Não interprete essa redução como melhora definitiva até o mês amadurecer.", False),
            ("", False),
            ("FONTES DE DADOS:", True),
            (f"Período analisado: {data_ini} a {data_fim}", False),
            ("SF2 (Notas de Saída) — relatório Gobi 595: nota fiscal, emissão e valor bruto.", False),
            ("SE1 (Contas a Receber) — relatório Gobi 596: título, parcela e vencimentos.", False),
            ("SE5 (Movimentação Bancária) — relatório Gobi 592: baixa, data e valor recebido.", False),
            ("O cálculo financeiro desta aba usa o encadeamento E5 → SE1 → SF2; não usa a dimensão interna de clientes como fonte de cálculo.", False),
            ("", False),
            ("NOTAS IMPORTANTES:", True),
            ("- Apenas títulos E1 liquidados por movimentos E5 e ligados a uma emissão SF2 entram no cálculo.", False),
            ("- Títulos ainda não pagos não possuem data de pagamento e ficam fora das três métricas.", False),
            ("- A aba Dados Brutos está no nível de parcela e movimento E5, não no nível de NF.", False),
            ("- O resumo consolida a Razão Social, enquanto o detalhe preserva cada CNPJ, loja, parcela e movimento bancário.", False),
            ("- Delta Atraso = PMR Pagamento − PMR Condição de Pagamento; positivo significa pagamento após o prazo contratado.", False),
            ("- O valor E5 utilizado é o valor recebido líquido de desconto, conforme a extração.", False),
            ("- A E1 é considerada somente quando o título possui emissão; sem emissão não há PMR.", False),
        ]
        for i, (txt, bold) in enumerate(texto, 1):
            c = ws4.cell(row=i, column=1, value=txt)
            if bold:
                c.font = titulo_font
        ws4.column_dimensions["A"].width = 115
        for row in ws4.iter_rows():
            cell = row[0]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value:
                cell_len = len(str(cell.value))
                ws4.row_dimensions[cell.row].height = 34 if cell_len > 140 else 22

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    try:
        conteudo = await asyncio.to_thread(_gerar)
        nome = f"PMR_{data_ini}_{data_fim}.xlsx"
        return StreamingResponse(
            io.BytesIO(conteudo),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{nome}"'},
        )
    except Exception as e:
        logger.exception("pmr_exportar: %s", e)
        raise HTTPException(500, f"Erro ao gerar Excel: {e}")
